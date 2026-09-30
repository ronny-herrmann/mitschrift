"""Aus einem ausgefüllten Heilbronner Word-Dokument (Briefkopf) eine leere Vorlage für den Protokollanten machen.

    python scripts/briefkopf_vorlage.py Eingabe.docx app/templates/briefkopf-heilbronn.docx

Behält Kopf-/Fußzeilen, Formatvorlagen, den Absenderblock und den grauen Titelkasten. Entfernt den Inhalt.
Felder werden zu Platzhaltern: {{AMT}} {{DATUM}} {{GZ}} {{TELEFON}} {{TITEL}}.
"""
import copy
import sys

from docx import Document
from docx.oxml.ns import qn


def set_text(p, text):
    """Absatz leeren (auch VIS-Feldfunktionen) und einen Lauf mit Text setzen – Zeichenformat des ersten Laufs bleibt."""
    rpr = None
    for r in p._p.findall(qn("w:r")):
        if r.find(qn("w:t")) is not None and r.find(qn("w:rPr")) is not None:
            rpr = copy.deepcopy(r.find(qn("w:rPr")))
            break
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    run = p.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def main(src, dst):
    d = Document(src)
    t0, t1 = d.tables[0], d.tables[1]
    set_text(t0.cell(0, 3).paragraphs[0], "{{DATUM}}")
    set_text(t0.cell(1, 0).paragraphs[0], "{{AMT}}")
    set_text(t0.cell(1, 3).paragraphs[0], "{{GZ}}")
    set_text(t0.cell(2, 3).paragraphs[0], "{{TELEFON}}")
    set_text(t1.cell(0, 0).paragraphs[0], "{{TITEL}}")
    body = d.element.body
    keep = {t0._tbl, t1._tbl}
    seen_t1 = False
    for el in list(body):
        if el.tag == qn("w:sectPr"):
            continue
        if el is t1._tbl:
            seen_t1 = True
            continue
        if el in keep:
            continue
        if not seen_t1 and el.tag == qn("w:p"):   # Abstand zwischen Absender und Titelkasten behalten
            continue
        body.remove(el)
    d.core_properties.title = "Protokoll"
    d.core_properties.subject = ""
    d.core_properties.comments = ""
    d.core_properties.author = "Stadt Heilbronn"
    d.core_properties.last_modified_by = "Protokollant"
    d.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
