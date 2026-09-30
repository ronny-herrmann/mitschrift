"""Protokoll-Modul: Prüfung, Zitatabgleich, Prompt, KI-Pipeline mit Fake-LLM."""

import json

from app.protokoll import LLMClient, bereinigt, build_prompt, check_quote, create_protokoll, verify
from app.store import Segment, Transcript


def make_transcript():
    t = Transcript(id="t1", title="Ausschuss", created_at="2026-09-26T10:00:00+00:00", source="upload", status="done", duration=90)
    t.segments = [
        Segment(0, 0.0, 4.0, "Guten Tag, wir beginnen mit der Sitzung.", "Vorsitz"),
        Segment(1, 5.0, 9.0, "Der erste Tagesordnungspunkt ist der Haushalt 2027.", "Vorsitz"),
        Segment(2, 10.0, 14.0, "Wir beschließen, den Haushaltsentwurf an den Gemeinderat weiterzuleiten.", "Vorsitz"),
        Segment(3, 15.0, 18.0, "Frau Müller übernimmt die Nachfrage bei der Kämmerei bis Freitag.", "Vorsitz"),
    ]
    return t


def test_verify_counts():
    t = make_transcript()
    md = "## Themen\n- Haushalt 2027 vorgestellt. [S1]\n- Etwas ohne Beleg.\n- Beleg auf falsches Segment. [S7]\n- Unverständlich [unklar]\n## Beschlüsse\nkeine"
    p = verify(t, md)
    assert p["zeilen_inhalt"] == 4
    assert p["unbelegt"] == 1
    assert p["ungueltige_belege"] == 1
    assert [z["status"] for z in p["zeilen"] if z["text"].strip() and not z["heading"] and z["text"].strip() != "keine"] == [
        "ok", "unbelegt", "ungueltig", "unklar"]


def test_verify_multi_refs():
    t = make_transcript()
    p = verify(t, "- Beschluss mit zwei Belegen. [S2, S3]\n- Anders geschrieben [S 1; S3]")
    assert [z["refs"] for z in p["zeilen"]] == [[2, 3], [1, 3]]
    assert p["unbelegt"] == 0 and p["ungueltige_belege"] == 0


def test_check_quote():
    seg = "Wir beschließen, den Haushaltsentwurf an den Gemeinderat weiterzuleiten."
    assert check_quote("den Haushaltsentwurf an den Gemeinderat", seg) == 1.0
    assert check_quote("Haushaltsentwurf an Gemeinderat weiterleiten", seg) > 0.6
    assert check_quote("Die Sitzung wird vertagt.", seg) < 0.6


def test_prompt_contains_rules_and_segments():
    p = build_prompt(make_transcript(), "verlauf")
    assert "Verlaufsprotokoll" in p and "[S3 | 00:15 | Vorsitz]" in p and "[unklar]" in p


class FakeLLM(LLMClient):
    """Liefert Zeilen im Extraktionsformat – eine davon frei erfunden."""

    def __init__(self, antwort=None):
        super().__init__("http://fake", "", "fake-model")
        self.calls = 0
        self.antwort = antwort

    def chat(self, system, user, temperature=0.0, json_mode=False):
        self.calls += 1
        if self.antwort is not None:
            return self.antwort
        out = []
        if "[S1 |" in user:
            out += ["THEMA: Haushalt 2027 [S1]", "AUSSAGE: Der erste Tagesordnungspunkt ist der Haushalt 2027. [S1]"]
        if "[S2 |" in user:
            out.append("BESCHLUSS: Der Haushaltsentwurf wird an den Gemeinderat weitergeleitet. [S2]")
        if "[S3 |" in user:
            out.append("AUFGABE: Frau Müller übernimmt die Nachfrage bei der Kämmerei bis Freitag. [S7]")  # falsche Nummer
            out.append("AUSSAGE: Die Personalplanung im Rahmen der Ressourcenhaushaltsplanung wird geprüft. [S3]")  # erfunden
        return "\n".join(out)


def test_create_protokoll_pipeline():
    llm = FakeLLM()
    c = create_protokoll(make_transcript(), "ergebnis", llm)
    assert llm.calls == 1
    md = c["protokoll_md"]
    assert "Personalplanung" not in md, "erfundene Aussage muss verworfen werden"
    assert c["verworfen"] == 1
    assert "[S2]" in md and "## Beschlüsse" in md
    assert "Frau Müller übernimmt" in md and "[S3]" in md, "falsche Nummer wird auf die passende Stelle korrigiert"
    assert c["pruefung"]["unbelegt"] == 0 and c["pruefung"]["ungueltige_belege"] == 0


def test_styles_differ_and_cache_skips_llm():
    t = make_transcript()
    llm = FakeLLM()
    first = create_protokoll(t, "zusammenfassung", llm)
    assert "## Überblick" in first["protokoll_md"] and "## Kernaussagen" in first["protokoll_md"]
    again = create_protokoll(t, "verlauf", None, cache=first["extrakt"])
    assert llm.calls == 1, "Wechsel der Protokollart braucht keine neue KI-Anfrage"
    assert "**Vorsitz:**" in again["protokoll_md"] and "00:00 – " in again["protokoll_md"]
    t.segments[0].clean = "Geändert"
    create_protokoll(t, "ergebnis", llm, cache=first["extrakt"])
    assert llm.calls == 2, "geänderter Text → neue Auswertung"


def test_everything_invented_gives_no_protokoll():
    llm = FakeLLM(antwort="THEMA: Rahmenbedingungen der Personalplanung [S1]\nAUSSAGE: Die Ressourcenhaushaltsplanung wird geprüft. [S2]")
    import pytest
    with pytest.raises(ValueError, match="nichts erfunden"):
        create_protokoll(make_transcript(), "zusammenfassung", llm)


def test_support():
    from app.protokoll import support
    assert support("Frau Müller übernimmt die Nachfrage bei der Kämmerei", "Frau Müller übernimmt die Nachfrage bei der Kämmerei bis Freitag.") == 1.0
    assert support("Personalplanung und Ressourcenhaushalt", "Winnie Puuh seufzte und ging nach Hause.") == 0.0


def test_bereinigt():
    assert bereinigt("ähm also wir, äh, beschließen das. Hm, genau.") == "Also wir, beschließen das. Genau."


# --- KI-Bereinigung ---------------------------------------------------------------
from app import bereinigung  # noqa: E402


def test_fidelity_rules():
    ok, _ = bereinigung.fidelity("ähm wir nutzen das CHS für die Tabellen", "Wir nutzen das Excel für die Tabellen.")
    assert ok
    ok, why = bereinigung.fidelity("Der Haushalt wird vorgestellt.", "Der Gemeinderat hat den Haushalt 2027 mit 12 Stimmen beschlossen und vertagt.")
    assert not ok, why
    ok, _ = bereinigung.fidelity("wir treffen uns um drei", "Wir treffen uns um 3.")
    assert ok
    ok, why = bereinigung.fidelity("wir treffen uns am montag", "Wir treffen uns am Montag um 14 Uhr.")
    assert not ok and "Zahlen" in why
    ok, why = bereinigung.fidelity("Ein Beinstand.", "Einbeinstand.")
    assert ok, why
    ok, why = bereinigung.fidelity("äh äh ja also wir wir machen das", "Ja, also wir machen das.")
    assert ok, why


def test_parse_and_prompt():
    t = make_transcript()
    p = bereinigung.build_prompt(t, [{"von": "CHS", "zu": "Excel"}])
    assert "[S0] Guten Tag" in p and "Excel" in p and "CHS → Excel" in p
    assert bereinigung.parse_lines("[S0] Hallo.\nMüll\n[ S 3 ] Test") == {0: "Hallo.", 3: "Test"}


class CleanLLM(LLMClient):
    def __init__(self):
        super().__init__("http://fake", "", "fake")

    def chat(self, system, user, temperature=0.0, json_mode=False):
        return "\n".join(line.replace("Haushalt 2027", "Haushalt 2027 (korrigiert)") if "[S1]" in line
                         else ("[S2] Etwas völlig anderes wurde hier frei erfunden und hinzugefügt." if "[S2]" in line else line)
                         for line in user.split("ZEILEN:\n", 1)[1].splitlines())


def test_clean_segments_with_fake_llm():
    t = make_transcript()
    cleaned = bereinigung.clean_segments(CleanLLM(), t.segments, [])
    changes, ok_n, bad_n = bereinigung.apply_cleaned(t.segments, cleaned)
    assert ok_n == 3 and bad_n == 1
    assert [c["idx"] for c in changes if not c["ok"]] == [2]


def test_create_protokoll_in_blocks():
    llm = FakeLLM()
    c = create_protokoll(make_transcript(), "ergebnis", llm, max_chars=120)
    assert llm.calls >= 3  # mehrere Abschnitte
    assert len(c["extrakt"]["abschnitte"]) == llm.calls


def test_korrektur_paare_und_weitere_stellen():
    assert bereinigung.korrektur_paare("und dann Bini Kum seufzte", "und dann Winnie Puuh seufzte") == [{"von": "Bini Kum", "zu": "Winnie Puuh"}]
    assert bereinigung.korrektur_paare("gleich", "Gleich") == []
    segs = [Segment(0, 0, 1, "Bini Kum seufzte"), Segment(1, 1, 2, "Wo ist bini kum?"), Segment(2, 2, 3, "Nichts")]
    assert bereinigung.weitere_stellen(segs, [{"von": "Bini Kum", "zu": "Winnie Puuh"}], 0) == [1]


def test_edited_segments_are_not_sent_and_corrections_win():
    t = make_transcript()
    t.segments[1].edited = True
    seen = []

    class LLM(LLMClient):
        def __init__(self):
            super().__init__("http://fake", "", "fake")

        def chat(self, system, user, temperature=0.0, json_mode=False):
            zeilen = user.split("ZEILEN:\n", 1)[1]
            seen.append(zeilen)
            return zeilen.replace("Kämmerei", "Bini Kum")

    kor = [{"von": "Bini Kum", "zu": "Kämmerei"}]
    out = bereinigung.clean_segments(LLM(), t.segments, [], korrekturen=kor)
    assert "[S1]" not in seen[0] and 1 not in out
    assert out[3].endswith("Kämmerei bis Freitag."), "Korrektur der Nutzer setzt sich gegen die KI durch"


def test_agenda_structures_result_protocol_and_tasks():
    t = make_transcript()
    antwort = ("THEMA: Haushalt [S1]\nAUSSAGE (TOP 1): Der erste Tagesordnungspunkt ist der Haushalt 2027. [S1]\n"
               "BESCHLUSS (TOP 1): Der Haushaltsentwurf wird an den Gemeinderat weitergeleitet. [S2]\n"
               "AUFGABE (TOP 2): Frau Müller | Nachfrage bei der Kämmerei | Freitag [S3]")
    c = create_protokoll(t, "ergebnis", FakeLLM(antwort=antwort), tagesordnung=["Haushalt 2027", "Kämmerei"])
    md = c["protokoll_md"]
    assert "## TOP 1: Haushalt 2027" in md and "## TOP 2: Kämmerei" in md
    assert "Frau Müller: Nachfrage bei der Kämmerei – bis Freitag" in md
    auf = [it for a in c["extrakt"]["abschnitte"] for it in a["items"] if it["typ"] == "AUFGABE"]
    assert auf[0]["wer"] == "Frau Müller" and auf[0]["bis"] == "Freitag"
    assert c["pruefung"]["unbelegt"] == 0


def test_agenda_from_word_table_with_letterhead():
    """Typische Besprechungsvorlage: Briefkopf-Tabelle + Tabelle mit Spalten Nr/Gemeldet von/Thema/Sachverhalt."""
    import io
    from docx import Document
    from app.glossar import read_agenda
    doc = Document()
    kopf = doc.add_table(rows=3, cols=4)
    for r, (a, b, c) in enumerate([("Stadt Musterstadt", "Datum", "01.01.2030"), ("Musteramt", "Gz.", "000"), ("", "Telefon", "123")]):
        kopf.cell(r, 0).text, kopf.cell(r, 1).text, kopf.cell(r, 3).text = a, b, c
    t = doc.add_table(rows=5, cols=6)
    for i, h in enumerate(["Lfd. Nr.", "Gemeldet von", "Thema", "Sachverhalt/Lösungsvorschlag", "Endergebnis", "Wer"]):
        t.cell(0, i).text = h
    t.cell(1, 1).text, t.cell(1, 2).text, t.cell(1, 3).text = "Ab", "Virtuelles Rathaus", "Ausweitung?\nÜbersetzung?"
    t.cell(2, 2).text = "Grippeimpfung"
    buf = io.BytesIO(); doc.save(buf)
    assert read_agenda("to.docx", buf.getvalue()) == ["Virtuelles Rathaus – Ausweitung? Übersetzung?", "Grippeimpfung"]


def test_agenda_from_text_prefers_numbered_lines():
    from app.glossar import read_agenda
    txt = "Stadt Musterstadt\nDatum 01.01.2030\nTagesordnung\nTOP 1: Begrüßung\nTOP 2: Haushalt\n3. Verschiedenes\n"
    assert read_agenda("to.txt", txt.encode()) == ["Begrüßung", "Haushalt", "Verschiedenes"]


def test_protokoll_laenge_waehlt_aus():
    from app.protokoll import _auswahl, assemble
    assert _auswahl(list(range(10)), 3) == [0, 4, 9] or len(_auswahl(list(range(10)), 3)) == 3
    woerter = ["Haushalt", "Personal", "Schulen", "Brücken", "Radwege", "Kitas", "Friedhof", "Bäder", "Parks"]
    items = [{"typ": "AUSSAGE", "text": f"Aussage über {w} im Detail", "refs": [0], "top": 1} for w in woerter]
    items.append({"typ": "BESCHLUSS", "text": "Beschluss A", "refs": [0], "top": 1})
    t = make_transcript()
    kurz = assemble(t, "ergebnis", [{"start": 0, "end": 1, "items": items}], ["Haushalt"], "kompakt")
    lang = assemble(t, "ergebnis", [{"start": 0, "end": 1, "items": items}], ["Haushalt"], "ausfuehrlich")
    assert kurz.count("Aussage ") == 2 and lang.count("Aussage ") == 9
    assert "Beschluss A" in kurz


def test_merge_short_speech_pieces(app_env):
    import numpy as np
    from app.pipeline import merge_short
    from app.vad import SpeechSegment
    sr = 16000
    audio = np.zeros(sr * 40, dtype=np.float32)
    seg = lambda a, b: SpeechSegment(int(a * sr), int(b * sr), audio[int(a * sr):int(b * sr)])  # noqa: E731
    # „Um“ / „Tonastaakin.“ / „Vierundzwierten“ → ein Stück, der lange Satz danach bleibt eigenständig
    out = merge_short([seg(3, 5), seg(6, 9), seg(12, 14), seg(15, 30), seg(36, 38)], audio)
    assert [(s.start // sr, s.end // sr) for s in out] == [(3, 14), (15, 30), (36, 38)]
