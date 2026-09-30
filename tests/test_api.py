"""API-Tests (Fake-Backend): Upload, Transkript, Bearbeitung, Export, Glossar, Live-WebSocket."""

import json
import time

import numpy as np
import soundfile as sf

from tests.conftest import FIXTURES


def wait_done(client, tid, timeout=30):
    t0 = time.time()
    while time.time() - t0 < timeout:
        d = client.get(f"/api/transcripts/{tid}").json()
        if d["status"] in ("done", "error"):
            return d
        time.sleep(0.2)
    raise AssertionError("Job nicht fertig")


def test_health(client):
    h = client.get("/api/health").json()
    assert h["model"]["backend"] == "fake"
    assert "ergebnis" in h["styles"]


def test_upload_flow(client):
    with open(FIXTURES / "drei_saetze_de.wav", "rb") as f:
        r = client.post("/api/upload", files={"file": ("drei_saetze_de.wav", f, "audio/wav")}, data={"title": "Testsitzung"})
    assert r.status_code == 200, r.text
    tid = r.json()["id"]
    d = wait_done(client, tid)
    assert d["status"] == "done", d
    assert d["title"] == "Testsitzung"
    assert 12 < d["duration"] < 13
    assert len(d["segments"]) == 3
    assert d["segments"][0]["start"] < d["segments"][1]["start"] < d["segments"][2]["start"]
    assert all(s["text"] for s in d["segments"])
    assert d["has_audio"]

    # Liste
    items = client.get("/api/transcripts").json()
    assert any(i["id"] == tid for i in items)

    # Bearbeiten
    r = client.patch(f"/api/transcripts/{tid}/segments/0", json={"text": "Korrigierter Text", "speaker": "Sprecher 1"})
    assert r.status_code == 200
    r = client.patch(f"/api/transcripts/{tid}/segments/1", json={"speaker": "Sprecher 1"})
    r = client.post(f"/api/transcripts/{tid}/speakers/rename", json={"von": "Sprecher 1", "zu": "Frau Müller"})
    assert r.json()["geaendert"] == 2
    d = client.get(f"/api/transcripts/{tid}").json()
    assert d["segments"][0]["text"] == "Korrigierter Text"
    assert d["segments"][0]["speaker"] == "Frau Müller"

    # Titel
    client.patch(f"/api/transcripts/{tid}", json={"title": "Neuer Titel"})
    assert client.get(f"/api/transcripts/{tid}").json()["title"] == "Neuer Titel"

    # Exporte
    txt = client.get(f"/api/transcripts/{tid}/export?format=txt").text
    assert "Frau Müller: Korrigierter Text" in txt and "[00:00]" in txt
    srt = client.get(f"/api/transcripts/{tid}/export?format=srt").text
    assert "-->" in srt
    md = client.get(f"/api/transcripts/{tid}/export?format=md").text
    assert md.startswith("# Neuer Titel")
    docx = client.get(f"/api/transcripts/{tid}/export?format=docx")
    assert docx.status_code == 200 and docx.content[:2] == b"PK"

    # NOVA-Prompt enthält Regeln + Segmente
    prompt = client.get(f"/api/transcripts/{tid}/nova-prompt?style=ergebnis").text
    assert "[S0 |" in prompt and "Ergebnisprotokoll" in prompt and "[unklar]" in prompt

    # Protokoll einfügen und prüfen lassen
    md_in = "## Ergebnisse\n- Der Haushalt wurde vorgestellt. [S1]\n- Ohne Beleg.\n## Beschlüsse\nkeine\n- Falscher Beleg [S99]"
    r = client.post(f"/api/transcripts/{tid}/protokoll/import", json={"style": "ergebnis", "protokoll_md": md_in})
    assert r.status_code == 200, r.text
    p = r.json()["content"]["pruefung"]
    assert p["zeilen_inhalt"] == 3 and p["unbelegt"] == 1 and p["ungueltige_belege"] == 1
    statuses = [z["status"] for z in p["zeilen"] if not z["heading"] and z["text"].strip()]
    # Der Testtext der Fake-Spracherkennung enthält nichts zum Haushalt → Beleg formal gültig, inhaltlich schwach
    assert statuses[:2] == ["schwach", "unbelegt"] and statuses[-1] == "ungueltig"

    # Bestätigen
    r = client.patch(f"/api/transcripts/{tid}/protokoll", json={"status": "bestaetigt"})
    assert r.json()["status"] == "bestaetigt"

    # Audio löschen, Transkript bleibt
    assert client.delete(f"/api/transcripts/{tid}/audio").status_code == 200
    d = client.get(f"/api/transcripts/{tid}").json()
    assert not d["has_audio"] and len(d["segments"]) == 3

    # Ganz löschen
    assert client.delete(f"/api/transcripts/{tid}").status_code == 200
    assert client.get(f"/api/transcripts/{tid}").status_code == 404


def test_upload_rejects_unknown_type(client):
    r = client.post("/api/upload", files={"file": ("x.exe", b"abc", "application/octet-stream")})
    assert r.status_code == 400


def test_glossar_applies(client):
    client.put("/api/glossar", json={"eintraege": [{"von": "Testtext", "zu": "Glossartext"}, {"von": "", "zu": "x"}]})
    assert client.get("/api/glossar").json() == [{"von": "Testtext", "zu": "Glossartext", "amt": ""}]
    with open(FIXTURES / "drei_saetze_de.wav", "rb") as f:
        tid = client.post("/api/upload", files={"file": ("a.wav", f, "audio/wav")}).json()["id"]
    d = wait_done(client, tid)
    assert all("Glossartext" in s["text"] for s in d["segments"])
    client.put("/api/glossar", json={"eintraege": []})
    client.delete(f"/api/transcripts/{tid}")


def test_live_websocket(client):
    x, sr = sf.read(FIXTURES / "drei_saetze_de.wav", dtype="float32")
    pcm = (np.clip(x, -1, 1) * 32767).astype("<i2").tobytes()
    frame = 3200  # 100 ms
    with client.websocket_connect("/ws/live?title=Livetest") as ws:
        ready = ws.receive_json()
        assert ready["type"] == "ready"
        tid = ready["transcript_id"]
        for i in range(0, len(pcm), frame):
            ws.send_bytes(pcm[i:i + frame])
        ws.send_text(json.dumps({"type": "stop"}))
        segments = []
        final = None
        while True:
            msg = ws.receive_json()
            if msg["type"] == "segment":
                segments.append(msg)
            elif msg["type"] == "final":
                final = msg
                break
            elif msg["type"] == "error":
                raise AssertionError(msg)
        assert len(segments) == 3, segments
        assert [s["idx"] for s in segments] == [0, 1, 2]
        assert final["transcript_id"] == tid
        assert 12 < final["duration"] < 13
        assert final["refining"] is True
        assert len(final["segments"]) == 3  # Live-Ergebnis sofort gespeichert
    d = wait_done(client, tid)  # Verfeinerung im Hintergrund
    assert d["status"] == "done" and d["source"] == "live" and d["has_audio"]
    assert len(d["segments"]) == 3
    audio = client.get(f"/api/transcripts/{tid}/audio")
    assert audio.status_code == 200 and audio.content[:4] == b"RIFF"
    client.delete(f"/api/transcripts/{tid}")


def test_live_disconnect_is_saved(client):
    x, sr = sf.read(FIXTURES / "drei_saetze_de.wav", dtype="float32")
    pcm = (np.clip(x, -1, 1) * 32767).astype("<i2").tobytes()
    with client.websocket_connect("/ws/live?title=Abbruch") as ws:
        tid = ws.receive_json()["transcript_id"]
        for i in range(0, len(pcm) // 2, 3200):
            ws.send_bytes(pcm[i:i + 3200])
        # Verbindung einfach schließen → Server sichert
    d = wait_done(client, tid)
    assert d["status"] == "done" and d["duration"] > 5
    client.delete(f"/api/transcripts/{tid}")


def test_login_required_when_password_set(client, monkeypatch):
    from app import auth
    from app.config import settings
    monkeypatch.setattr(settings, "access_password", "geheim-123")
    client.cookies.clear()
    assert client.get("/api/transcripts").status_code == 401
    r = client.get("/", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"
    assert client.post("/api/login", json={"password": "falsch"}).status_code == 401
    r = client.post("/api/login", json={"password": "geheim-123"})
    assert r.status_code == 204 and auth.COOKIE in r.cookies
    client.cookies.set(auth.COOKIE, r.cookies[auth.COOKIE])
    assert client.get("/api/transcripts").status_code == 200
    with client.websocket_connect("/ws/live") as ws:
        assert ws.receive_json()["type"] == "ready"
        ws.send_text(json.dumps({"type": "stop"}))
        while ws.receive_json()["type"] != "final":
            pass
    client.cookies.clear()
    import pytest
    from starlette.websockets import WebSocketDisconnect
    with pytest.raises(WebSocketDisconnect):
        with client.websocket_connect("/ws/live") as ws:
            ws.receive_json()
    auth._failures.clear()


def test_bereinigen_import_and_discard(client):
    with open(FIXTURES / "drei_saetze_de.wav", "rb") as f:
        tid = client.post("/api/upload", files={"file": ("a.wav", f, "audio/wav")}).json()["id"]
    d = wait_done(client, tid)
    prompt = client.get(f"/api/transcripts/{tid}/bereinigen-prompt").text
    assert "[S0]" in prompt and "Verboten" in prompt
    assert client.post(f"/api/transcripts/{tid}/bereinigen").status_code == 409  # keine KI angebunden
    orig = d["segments"][0]["text"]
    r = client.post(f"/api/transcripts/{tid}/bereinigen/import",
                    json={"text": f"[S0] {orig.strip('[]')}.\n[S1] Der Bürgermeister hat heute 5 Millionen Euro für das neue Rathaus beschlossen."})
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["uebernommen"] == 1 and res["verworfen"] == 1
    d = client.get(f"/api/transcripts/{tid}").json()
    assert d["segments"][0]["clean"] and not d["segments"][1]["clean"] and d["segments"][1]["clean_note"]
    txt = client.get(f"/api/transcripts/{tid}/export?format=txt").text
    assert d["segments"][0]["clean"] in txt
    client.delete(f"/api/transcripts/{tid}/bereinigen")
    assert not client.get(f"/api/transcripts/{tid}").json()["segments"][0]["clean"]
    client.delete(f"/api/transcripts/{tid}")


def test_docs_endpoint(client):
    r = client.get("/api/docs/faq")
    assert r.status_code == 200 and "<details>" in r.text
    assert client.get("/api/docs/../main").status_code == 404
    assert client.get("/api/docs/unbekannt").status_code == 404
    assert "script-src 'self'" in r.headers["content-security-policy"]


def _upload(client, title="Korrekturtest"):
    with open(FIXTURES / "drei_saetze_de.wav", "rb") as f:
        tid = client.post("/api/upload", files={"file": ("a.wav", f, "audio/wav")}, data={"title": title}).json()["id"]
    wait_done(client, tid)
    return tid


def test_correction_priority_and_replace_elsewhere(client):
    tid = _upload(client)
    client.patch(f"/api/transcripts/{tid}/segments/1", json={"text": "Bini Kum seufzte"})
    client.patch(f"/api/transcripts/{tid}/segments/2", json={"text": "Wo ist Bini Kum?"})
    r = client.patch(f"/api/transcripts/{tid}/segments/1", json={"text": "Winnie Puuh seufzte"}).json()
    assert r["korrekturen"] == [{"von": "Bini Kum", "zu": "Winnie Puuh"}]
    assert r["weitere_stellen"] == [2]
    assert client.post(f"/api/transcripts/{tid}/ersetzen", json={"paare": r["korrekturen"]}).json()["ersetzt"] == 1
    d = client.get(f"/api/transcripts/{tid}").json()
    assert d["segments"][2]["text"] == "Wo ist Winnie Puuh?" and d["segments"][1]["edited"] and d["segments"][2]["edited"]
    assert "Winnie Puuh" in client.get(f"/api/transcripts/{tid}/bereinigen-prompt").text


def test_export_names_and_fliesstext(client):
    tid = _upload(client, "Test RH")
    r = client.get(f"/api/transcripts/{tid}/export?format=fliesstext")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert 'filename="Test RH - Fliesstext.docx"' in r.headers["content-disposition"]
    r = client.get(f"/api/transcripts/{tid}/export?format=md")
    assert "Test%20RH%20-%20Markdown.md" in r.headers["content-disposition"]
    assert client.get(f"/api/transcripts/{tid}/export?format=docx&protokoll=1").status_code == 404


def test_ai_tasks_run_in_background(client, monkeypatch):
    import time as _t
    from app import main
    from app.protokoll import LLMClient

    class LLM(LLMClient):
        def __init__(self):
            super().__init__("http://fake", "", "fake")

        def chat(self, system, user, temperature=0.0, json_mode=False):
            if "ZEILEN:\n" in user:
                return user.split("ZEILEN:\n", 1)[1]
            nums = __import__("re").findall(r"\[S(\d+) \|[^\]]*\] (.*)", user)
            return "\n".join(f"AUSSAGE: {txt.strip('[]')} [S{n}]" for n, txt in nums) + "\nAUSSAGE: Erfundener Haushaltsbeschluss [S0]"

    monkeypatch.setattr(main, "get_llm", lambda: LLM())
    monkeypatch.setattr(main.settings, "llm_base_url", "http://fake")
    monkeypatch.setattr(main.settings, "llm_model", "fake")
    tid = _upload(client)

    def wait(tid):
        for _ in range(100):
            s = client.get(f"/api/transcripts/{tid}/task").json()
            if s["status"] != "running":
                return s
            _t.sleep(0.05)
        raise AssertionError("Aufgabe hängt")

    assert client.post(f"/api/transcripts/{tid}/bereinigen").json()["started"]
    s = wait(tid)
    assert s["status"] == "done" and s["result"]["uebernommen"] == 3 and s["total"] >= 1
    assert client.post(f"/api/transcripts/{tid}/protokoll", json={"style": "verlauf"}).json()["started"]
    assert wait(tid)["status"] == "done"
    p = client.get(f"/api/transcripts/{tid}").json()["protokoll"]
    assert p["style"] == "verlauf" and "Erfundener" not in p["content"]["protokoll_md"]
    assert p["content"]["verworfen"] == 1


def test_several_protocols_and_protocol_export(client):
    tid = _upload(client, "Mehrere")
    for style in ("zusammenfassung", "ergebnis", "zusammenfassung"):
        r = client.post(f"/api/transcripts/{tid}/protokoll/import", json={"style": style, "protokoll_md": f"## {style}\n- Punkt [S0]"})
        assert r.status_code == 200
    d = client.get(f"/api/transcripts/{tid}").json()
    assert set(d["protokolle"]) == {"zusammenfassung", "ergebnis"}, "je Art genau ein Protokoll"
    r = client.patch(f"/api/transcripts/{tid}/protokoll", json={"style": "ergebnis", "status": "bestaetigt"}).json()
    assert r["status"] == "bestaetigt" and r["content"]["geprueft_am"]
    r = client.get(f"/api/transcripts/{tid}/export?format=protokoll&style=ergebnis")
    assert r.status_code == 200 and "Ergebnisprotokoll.docx" in r.headers["content-disposition"]
    from docx import Document
    import io
    paras = Document(io.BytesIO(r.content)).paragraphs
    text = "\n".join(p.text for p in paras)
    assert "Inhaltlich geprüft am" in text and "[S0]" not in text
    assert "Transkript" not in [p.text for p in paras if p.style.name.startswith("Heading")], "nur das Protokoll"


def test_glossar_per_amt_and_import(client):
    client.put("/api/glossar", json={"eintraege": [
        {"von": "Testtext", "zu": "Amt-A-Text", "amt": "10.5"},
        {"von": "Sekunden", "zu": "Sek.", "amt": ""},
    ]})
    assert client.get("/api/glossar/aemter").json() == ["10.5"]
    # Aufnahme für Amt 62.2: nur allgemeine Einträge greifen
    with open(FIXTURES / "drei_saetze_de.wav", "rb") as f:
        tid = client.post("/api/upload", files={"file": ("a.wav", f, "audio/wav")},
                          data={"glossar": "62.2", "teilnehmende": "Frau Müller, Herr Maier",
                                "tagesordnung": "TOP 1: Haushalt\n2. Personal"}).json()["id"]
    d = wait_done(client, tid)
    assert all("Amt-A-Text" not in s["text"] and "Sek." in s["text"] for s in d["segments"])
    assert d["meta"]["teilnehmende"] == ["Frau Müller", "Herr Maier"]
    assert d["meta"]["tagesordnung"] == ["Haushalt", "Personal"]
    # Excel-Import
    import io
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["Erkannt als", "Richtig", "Amt"]); ws.append(["CHS", "Excel", "20.1"]); ws.append(["Heilbrunn", "Heilbronn", None])
    buf = io.BytesIO(); wb.save(buf)
    r = client.post("/api/glossar/import", files={"file": ("g.xlsx", buf.getvalue(), "application/octet-stream")}, data={"amt": "62.2"})
    assert r.json() == [{"von": "CHS", "zu": "Excel", "amt": "20.1"}, {"von": "Heilbrunn", "zu": "Heilbronn", "amt": "62.2"}]
    # Tagesordnung aus Word
    from docx import Document
    doc = Document(); doc.add_paragraph("Tagesordnung"); doc.add_paragraph("TOP 1: Begrüßung"); doc.add_paragraph("TOP 2: Haushalt 2027")
    b2 = io.BytesIO(); doc.save(b2)
    punkte = client.post("/api/tagesordnung/lesen", files={"file": ("to.docx", b2.getvalue(), "application/octet-stream")}).json()
    assert punkte[-2:] == ["Begrüßung", "Haushalt 2027"]
    client.put("/api/glossar", json={"eintraege": []})


def test_change_access_password(client, monkeypatch):
    from app import auth, main
    from app.config import settings
    monkeypatch.setattr(settings, "access_password", "alt-passwort-1")
    client.cookies.clear()
    try:
        r = client.post("/api/login", json={"password": "alt-passwort-1"})
        client.cookies.set(auth.COOKIE, r.cookies[auth.COOKIE])
        alt_cookie = r.cookies[auth.COOKIE]
        assert client.post("/api/zugang", json={"alt": "falsch", "neu": "neues-passwort-2"}).status_code == 403
        assert client.post("/api/zugang", json={"alt": "alt-passwort-1", "neu": "kurz"}).status_code == 400
        r = client.post("/api/zugang", json={"alt": "alt-passwort-1", "neu": "neues-passwort-2"})
        assert r.status_code == 204
        client.cookies.set(auth.COOKIE, r.cookies[auth.COOKIE])
        assert client.get("/api/transcripts").status_code == 200
        # alte Sitzungen sind abgemeldet, altes Passwort gilt nicht mehr
        client.cookies.set(auth.COOKIE, alt_cookie)
        assert client.get("/api/transcripts").status_code == 401
        assert client.post("/api/login", json={"password": "alt-passwort-1"}).status_code == 401
        assert client.post("/api/login", json={"password": "neues-passwort-2"}).status_code == 204
    finally:
        main.state.store.set_setting("zugang_passwort", None)
        client.cookies.clear()
        auth._failures.clear()


def test_word_export_uses_letterhead(client):
    import io
    import docx
    with open(FIXTURES / "drei_saetze_de.wav", "rb") as f:
        tid = client.post("/api/upload", files={"file": ("a.wav", f, "audio/wav")}, data={"title": "Briefkopf-Test"}).json()["id"]
    wait_done(client, tid)
    r = client.get(f"/api/transcripts/{tid}/export", params={"format": "docx"})
    assert r.status_code == 200
    d = docx.Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for tb in d.tables for row in tb.rows for c in row.cells)
    assert "Briefkopf-Test" in text
    assert "{{" not in text
    assert "Herrmann" not in text and "Smartphone" not in text
    assert d.sections[0].header is not None


def test_audio_download(client):
    with open(FIXTURES / "drei_saetze_de.wav", "rb") as f:
        tid = client.post("/api/upload", files={"file": ("a.wav", f, "audio/wav")}, data={"title": "Audio-Test"}).json()["id"]
    wait_done(client, tid)
    r = client.get(f"/api/transcripts/{tid}/audio", params={"download": 1})
    assert r.status_code == 200 and r.content[:4] == b"RIFF"
    assert "attachment" in r.headers["content-disposition"] and "Audio-Test - Audio.wav" in r.headers["content-disposition"]
    assert "attachment" not in client.get(f"/api/transcripts/{tid}/audio").headers.get("content-disposition", "")
