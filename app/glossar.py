"""Glossar: Nachkorrektur bekannter Begriffe (Ämter, Gremien, Namen, Fachwörter).

Modelle erkennen Eigennamen nicht immer richtig. Das Glossar ersetzt bekannte
Fehlschreibungen wortgenau (Groß-/Kleinschreibung egal) – einfach, transparent
und ohne Modelltraining. Beispiel: "Heilbrunn" → "Heilbronn".
"""

from __future__ import annotations

import re


def compile_glossar(entries: list[dict]) -> list[tuple[re.Pattern, str]]:
    rules = []
    for e in entries or []:
        von = (e.get("von") or "").strip()
        zu = (e.get("zu") or "").strip()
        if not von or not zu or von.lower() == zu.lower():
            continue
        rules.append((re.compile(r"(?<!\w)" + re.escape(von) + r"(?!\w)", re.IGNORECASE), zu))
    return rules


def apply_glossar(text: str, rules: list[tuple[re.Pattern, str]]) -> str:
    for pat, zu in rules:
        text = pat.sub(zu, text)
    return text


# ---------------------------------------------------------------------------
# Import: Glossar aus Excel/CSV, Tagesordnung aus Word/PDF/Text
# ---------------------------------------------------------------------------
_HEAD_VON = ("erkannt", "falsch", "von", "fehl", "gehört", "statt")
_HEAD_ZU = ("richtig", "zu", "korrekt", "schreibweise", "soll")
_HEAD_AMT = ("amt", "bereich", "abteilung")


def _rows_to_entries(rows: list[list[str]]) -> list[dict]:
    rows = [[str(c or "").strip() for c in r] for r in rows if any(str(c or "").strip() for c in r)]
    if not rows:
        return []
    head = [c.lower() for c in rows[0]]

    def find(keys):
        for i, h in enumerate(head):
            if any(k in h for k in keys):
                return i
        return None

    iv, iz, ia = find(_HEAD_VON), find(_HEAD_ZU), find(_HEAD_AMT)
    if iv is not None and iz is not None and iv != iz:
        body = rows[1:]
    else:
        iv, iz, ia, body = 0, 1, (2 if len(rows[0]) > 2 else None), rows
    out = []
    for r in body:
        get = lambda i: r[i] if i is not None and i < len(r) else ""  # noqa: E731
        out.append({"von": get(iv), "zu": get(iz), "amt": get(ia)})
    return out


def read_table(filename: str, data: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return _rows_to_entries([list(r) for r in ws.iter_rows(values_only=True)])
    if name.endswith((".csv", ".txt")):
        import csv
        import io
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
        return _rows_to_entries(list(csv.reader(io.StringIO(text), dialect)))
    raise ValueError("Bitte eine Excel-Datei (.xlsx) oder CSV-Datei hochladen.")


def read_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    import io
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for tbl in doc.tables:
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" ".join(dict.fromkeys(cells)))
        return "\n".join(lines)
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages[:10])
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8-sig", errors="replace")
    raise ValueError("Bitte eine Word-Datei (.docx), PDF oder Textdatei hochladen.")


# ---------------------------------------------------------------------------
# Tagesordnung aus Datei lesen
# ---------------------------------------------------------------------------
_TOP_PREFIX = re.compile(r"^\s*(?:TOP\s*\d+[a-z]?|\d{1,2}(?:\.\d{1,2})*[.)]|[-*•–])\s*[:.)-]?\s*", re.I)
_NOISE = re.compile(
    r"^(?:stadt\s+\w+|datum|gz\.?|aktenzeichen|az\.?|telefon|tel\.?|telefax|fax|e-?mail|zimmer|raum|ort|zeit|uhrzeit|beginn|ende"
    r"|tagesordnung|einladung|protokoll|niederschrift|sitzung|besprechung|sehr geehrte.*|mit freundlichen.*|anlage.*|seite \d+.*"
    r"|[\w\- ]*amt|personal- und organisationsamt)\b[:\s]*[\w./ ,:-]{0,40}$", re.I)
_DATE_ONLY = re.compile(r"^[\d.\s:/-]+(?:uhr)?$", re.I)
_THEMA_KEYS = ("thema", "tagesordnungspunkt", "beratungsgegenstand", "gegenstand", "betreff", "titel", "punkt", "inhalt")
_DETAIL_KEYS = ("sachverhalt", "lösungsvorschlag", "loesungsvorschlag", "beschreibung", "erläuterung", "erlaeuterung", "fragestellung")
_KOPF_KEYS = ("datum", "gz.", "gz", "telefon", "aktenzeichen", "e-mail", "telefax", "zimmer")


def _clean_cell(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\n", " ")).strip(" -–")


def _row_cells(row) -> list[str]:
    """Zellen einer Tabellenzeile – verbundene Zellen nur einmal."""
    seen, out = set(), []
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        out.append(_clean_cell(c.text))
    return out


def _kurz(text: str, n: int = 160) -> str:
    return text if len(text) <= n else text[:n].rsplit(" ", 1)[0] + " …"


def _agenda_from_tables(tables) -> list[str]:
    for tbl in tables:
        rows = [_row_cells(r) for r in tbl.rows]
        flat = " ".join(" ".join(r) for r in rows).lower()
        if not rows or sum(k in flat.split() or k in flat for k in _KOPF_KEYS) >= 2 and len(rows) <= 4:
            continue  # Briefkopf (Datum, Gz., Telefon …)
        for hi, head in enumerate(rows[:3]):
            low = [h.lower() for h in head]
            col = next((i for k in _THEMA_KEYS for i, h in enumerate(low) if h.startswith(k) or f" {k}" in f" {h}"), None)
            if col is None:
                continue
            det = next((i for k in _DETAIL_KEYS for i, h in enumerate(low) if k in h and i != col), None)
            out = []
            for r in rows[hi + 1:]:
                if col >= len(r) or not r[col]:
                    continue
                item = r[col]
                if det is not None and det < len(r) and r[det]:
                    item = f"{item} – {r[det]}"
                out.append(_kurz(item))
            if out:
                return out
        # ohne Kopfzeile: „1. | Thema | …“ → Thema
        out = []
        for r in rows:
            if len(r) >= 2 and re.fullmatch(r"(?:TOP\s*)?\d{1,2}[.)]?", r[0], re.I):
                rest = next((c for c in r[1:] if c), "")
                if rest:
                    out.append(_kurz(rest))
        if len(out) >= 2:
            return out
    return []


def _agenda_from_lines(lines: list[str], numbered: set[int] | None = None) -> list[str]:
    lines = [_clean_cell(x) for x in lines]
    marked = [i for i, x in enumerate(lines) if x and (_TOP_PREFIX.match(x) or (numbered and i in numbered))]
    if len(marked) >= 2:
        pick = marked
    else:
        pick = [i for i, x in enumerate(lines) if x and not _NOISE.match(x) and not _DATE_ONLY.match(x)]
    out = []
    for i in pick:
        x = _TOP_PREFIX.sub("", lines[i], count=1).strip()
        if len(x) >= 2 and not _DATE_ONLY.match(x):
            out.append(_kurz(x))
    return out


def read_agenda(filename: str, data: bytes) -> list[str]:
    """Tagesordnungspunkte aus Word, PDF oder Text. Word-Tabellen mit Spalte „Thema“ (typische Heilbronner
    Besprechungsvorlage) werden erkannt; Briefkopf-Angaben (Stadt, Amt, Datum, Gz., Telefon) werden ignoriert."""
    import io
    name = filename.lower()
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        found = _agenda_from_tables(doc.tables)
        if found:
            return found
        paras = doc.paragraphs
        numbered = {i for i, p in enumerate(paras)
                    if p._p.pPr is not None and p._p.pPr.numPr is not None or (p.style is not None and p.style.name.lower().startswith("list"))}
        return _agenda_from_lines([p.text for p in paras], numbered)
    return _agenda_from_lines(read_text(filename, data).splitlines())
