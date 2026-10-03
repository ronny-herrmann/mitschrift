"""Export: TXT, Markdown, SRT (Untertitel), DOCX (Word)."""

from __future__ import annotations

import io
from datetime import datetime

from .store import Transcript


def fmt_dauer(sec: float) -> str:
    """1:09 Min. bzw. 1:02:03 Std. – für Menschen lesbar."""
    sec = int(round(sec or 0))
    h, rest = divmod(sec, 3600)
    m, s = divmod(rest, 60)
    return f"{h}:{m:02d}:{s:02d} Std." if h else f"{m}:{s:02d} Min."


def local_dt(iso: str) -> str:
    """Zeitstempel (UTC gespeichert) in deutscher Ortszeit anzeigen."""
    from zoneinfo import ZoneInfo
    try:
        return datetime.fromisoformat(iso).astimezone(ZoneInfo("Europe/Berlin")).strftime("%d.%m.%Y %H:%M")
    except Exception:
        return ""


def fmt_time(sec: float) -> str:
    sec = max(0.0, float(sec))
    h, rest = divmod(int(sec), 3600)
    m, s = divmod(rest, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _fmt_srt(sec: float) -> str:
    ms = int(round(sec * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def _line(seg, with_time: bool, with_speaker: bool) -> str:
    parts = []
    if with_time:
        parts.append(f"[{fmt_time(seg.start)}]")
    if with_speaker and seg.speaker:
        parts.append(f"{seg.speaker}:")
    parts.append(seg.clean or seg.text)
    return " ".join(parts)


def to_txt(t: Transcript, with_time: bool = True, with_speaker: bool = True) -> str:
    lines = [_line(s, with_time, with_speaker) for s in t.segments]
    return "\n".join(lines) + "\n"


def to_markdown(t: Transcript) -> str:
    created = local_dt(t.created_at) if t.created_at else ""
    out = [f"# {t.title}", "", f"Aufgenommen: {created} · Dauer: {fmt_time(t.duration)} · Modell: {t.model}", ""]
    for s in t.segments:
        out.append(f"- **{fmt_time(s.start)}** {('*' + s.speaker + ':* ') if s.speaker else ''}{s.clean or s.text}")
    return "\n".join(out) + "\n"


def to_srt(t: Transcript) -> str:
    out = []
    for i, s in enumerate(t.segments, 1):
        body = s.clean or s.text
        text = f"{s.speaker}: {body}" if s.speaker else body
        out += [str(i), f"{_fmt_srt(s.start)} --> {_fmt_srt(s.end)}", text, ""]
    return "\n".join(out)


def refs_to_times(md: str, t: Transcript) -> str:
    """[S12] → [03:41] – im Export sind Zeitmarken verständlicher als Segmentnummern."""
    import re
    starts = {s.idx: s.start for s in t.segments}

    def repl(m):
        nums = [int(n) for n in re.findall(r"\d+", m.group(0))]
        times = [fmt_time(starts[n]) for n in nums if n in starts]
        return f"[{', '.join(dict.fromkeys(times))}]" if times else ""
    return re.sub(r"\[\s*S\s*\d+(?:\s*[,;–-]\s*S?\s*\d+)*\s*\]", repl, md)


def strip_refs(md: str) -> str:
    """Belegnummern ([S12]) für Leser entfernen – die Prüfung ist dann schon erfolgt."""
    import re
    md = re.sub(r"\s*\[\s*S\s*\d+(?:\s*[,;–-]\s*S?\s*\d+)*\s*\]", "", md)
    return re.sub(r"[ \t]+\n", "\n", md)


def _md_inline(p, text: str) -> None:
    """**fett** aus dem Protokoll als fette Schrift übernehmen."""
    import re
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            p.add_run(part[2:-2]).bold = True
        elif part:
            p.add_run(part)


VORLAGEN_DIR = __import__("pathlib").Path(__file__).parent / "templates"


def vorlagen() -> list[str]:
    """Mitgelieferte Word-Briefköpfe (Dateinamen ohne Endung)."""
    return sorted(p.stem for p in VORLAGEN_DIR.glob("*.docx")) if VORLAGEN_DIR.exists() else []


EIGENE_DIR = None   # wird von main.py gesetzt: DATA_DIR/vorlagen (eigene, hochgeladene Briefköpfe)


def eigene_vorlagen() -> list[dict]:
    """Hochgeladene Briefköpfe: [{id, name}] – liegen auf dem Server, nicht im Quellcode."""
    import json
    if not EIGENE_DIR or not EIGENE_DIR.exists():
        return []
    out = []
    for f in sorted(EIGENE_DIR.glob("*.docx")):
        meta = f.with_suffix(".json")
        name = f.stem
        if meta.exists():
            try:
                name = json.loads(meta.read_text("utf-8")).get("name") or name
            except Exception:
                pass
        out.append({"id": f.stem, "name": name})
    return out


def _vorlage_pfad(vorlage: str | None):
    """„eigen:<id>“ → hochgeladene Datei, sonst mitgelieferte Vorlage (Standard: Stadt Heilbronn)."""
    if vorlage and vorlage.startswith("eigen:") and EIGENE_DIR:
        f = EIGENE_DIR / f"{vorlage[6:]}.docx"
        if f.exists() and f.parent == EIGENE_DIR:
            return f, True
    name = vorlage if vorlage in vorlagen() else (vorlagen()[0] if vorlagen() else None)
    return (VORLAGEN_DIR / f"{name}.docx", False) if name else (None, False)


def _hat_marker(doc) -> bool:
    texte = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return any("{{" in x for x in texte)


def _body_leeren(doc) -> None:
    """Inhalt entfernen, Kopf-/Fußzeilen, Seitenränder und Formatvorlagen bleiben."""
    from docx.oxml.ns import qn
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def _sdt(paragraph, placeholder: str, tag: str) -> None:
    """Ausfüllbares Feld (Word-Inhaltssteuerelement) mit grauem Platzhaltertext einfügen."""
    from docx.oxml import parse_xml
    xml = ('<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f'<w:sdtPr><w:alias w:val="{placeholder}"/><w:tag w:val="{tag}"/><w:showingPlcHdr/><w:text/></w:sdtPr>'
           f'<w:sdtContent><w:r><w:rPr><w:color w:val="808080"/></w:rPr><w:t xml:space="preserve">{placeholder}</w:t></w:r>'
           '</w:sdtContent></w:sdt>')
    paragraph._p.append(parse_xml(xml))


def _alle_absaetze(doc):
    """Absätze im Text, in Tabellen sowie in Kopf- und Fußzeilen."""
    def aus(container):
        for p in container.paragraphs:
            yield p
        for tbl in container.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    yield from aus(cell)
    yield from aus(doc)
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
            try:
                yield from aus(part)
            except Exception:
                pass


def _fill_marker(doc, marker: str, text: str | None = None, placeholder: str = "", tag: str = "") -> None:
    """Platzhalter {{…}} der Vorlage ersetzen – durch Text oder ein ausfüllbares Feld."""
    for p in _alle_absaetze(doc):
        if marker not in p.text:
            continue
        runs = p.runs
        if marker not in "".join(r.text for r in runs) or not runs:
            continue
        if text is not None:
            full = "".join(r.text for r in runs).replace(marker, text)
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            lines = full.split("\n")
            runs[0].text = lines[0]
            for extra in lines[1:]:
                runs[0].add_break()
                runs[0].add_text(extra)
        else:
            rest = "".join(r.text for r in runs).replace(marker, "")
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            runs[0].text = rest
            _sdt(p, placeholder, tag)


def _base_document(t: Transcript, titel: str, vorlage: str | None):
    """Neues Dokument auf Basis eines Briefkopfs. Rückgabe: (Dokument, Titelkasten vorhanden?)."""
    from docx import Document
    from docx.shared import Pt
    pfad, eigen = _vorlage_pfad(vorlage)
    if pfad:
        doc = Document(str(pfad))
        if not eigen or _hat_marker(doc):
            hat_titel = any("{{TITEL}}" in par.text for par in _alle_absaetze(doc))
            datum = local_dt(t.created_at)[:10] if t.created_at else ""
            _fill_marker(doc, "{{DATUM}}", datum)
            _fill_marker(doc, "{{TITEL}}", titel)
            _fill_marker(doc, "{{AMT}}", placeholder="Amt eintragen", tag="amt")
            _fill_marker(doc, "{{GZ}}", placeholder="Gz. eintragen", tag="gz")
            _fill_marker(doc, "{{TELEFON}}", placeholder="Telefon eintragen", tag="telefon")
            if hat_titel:
                return doc, True
        else:
            # eigener Briefkopf ohne Platzhalter: nur Kopf-/Fußzeile, Ränder und Schriften übernehmen
            _body_leeren(doc)
        p = doc.add_paragraph()
        r = p.add_run(titel.replace("\n", " – "))
        r.bold = True
        r.font.size = Pt(16)
        return doc, False
    doc = Document()
    doc.styles["Normal"].font.name = "Source Sans Pro"
    doc.styles["Normal"].font.size = Pt(11)
    _heading(doc, titel.replace("\n", " – "), 1)
    return doc, False


def _style_da(doc, name: str) -> bool:
    try:
        doc.styles[name]
        return True
    except KeyError:
        return False


def _heading(doc, text: str, level: int):
    """Überschrift – auch in fremden Vorlagen ohne Überschrift-Formatvorlagen."""
    from docx.shared import Pt
    if _style_da(doc, f"Heading {level}"):
        return doc.add_heading(text, level=level)
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt({1: 16, 2: 14, 3: 12.5}.get(level, 11.5))
    p.paragraph_format.space_before = Pt(10)
    return p


def _bullet(doc, text: str, briefkopf: bool):
    from docx.shared import Cm
    if not briefkopf and _style_da(doc, "List Bullet"):
        p = doc.add_paragraph(style="List Bullet")
    else:
        p = doc.add_paragraph(style="List Paragraph") if _style_da(doc, "List Paragraph") else doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.add_run("•\u00a0\u00a0")
    _md_inline(p, text)
    return p


def to_docx(t: Transcript, protokoll_md: str | None = None, protokoll_name: str = "Zusammenfassung",
            fliesstext: bool = False, nur_protokoll: bool = False, pruefstatus: str = "",
            teilnehmende: list[str] | None = None, vorlage: str | None = None) -> bytes:
    from docx.shared import Cm, Pt, RGBColor

    art = protokoll_name if protokoll_md else ("Fließtext" if fliesstext else "Transkript")
    doc, briefkopf = _base_document(t, f"{art}\n{t.title}", vorlage)

    # Sitzungsdaten als Kopfblock (Protokollführung ist ausfüllbar)
    created = local_dt(t.created_at) if t.created_at else ""
    if briefkopf:
        doc.add_paragraph()  # Abstand zum Titelkasten
    zeilen = [("Sitzung", t.title), ("Datum und Uhrzeit", created), ("Dauer", fmt_dauer(t.duration))]
    if teilnehmende:
        zeilen.append(("Teilnehmende", ", ".join(teilnehmende)))
    for label, wert in zeilen + [("Protokollführung", None)]:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.5))
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f"{label}:\t").bold = True
        if wert is None:
            _sdt(p, "Name eintragen", "protokollfuehrung")
        else:
            p.add_run(wert)
    hinweis = doc.add_paragraph(f"Automatisch erstellt mit dem Protokollanten der Stadt Heilbronn · Spracherkennung: {t.model}")
    hinweis.runs[0].font.size = Pt(8)
    hinweis.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    if fliesstext:
        # Absätze je Sprecherwechsel, ohne Zeitmarken
        cur_sp, buf = None, []
        for s in t.segments:
            if s.speaker != cur_sp and buf:
                doc.add_paragraph(" ".join(buf))
                buf = []
            cur_sp = s.speaker
            buf.append(s.clean or s.text)
        if buf:
            doc.add_paragraph(" ".join(buf))
    else:
        if protokoll_md:
            if not briefkopf:  # im Briefkopf steht die Art schon im Titelkasten
                _heading(doc, protokoll_name, 2)
            note = doc.add_paragraph((pruefstatus + ". " if pruefstatus else "")
                                     + "Jede Aussage wurde automatisch gegen das Transkript geprüft.")
            note.runs[0].font.size = Pt(9)
            note.runs[0].italic = True
            lines = [ln.rstrip() for ln in strip_refs(protokoll_md).splitlines()]
            section, i = "", 0
            while i < len(lines):
                line = lines[i]
                i += 1
                if not line:
                    continue
                # Aufgaben „Wer | Was | Bis wann“ als Tabelle
                if section.lower().startswith("aufgaben") and line.startswith(("- ", "* ")) and line.count("|") >= 2:
                    rows = [line]
                    while i < len(lines) and lines[i].startswith(("- ", "* ")) and lines[i].count("|") >= 2:
                        rows.append(lines[i])
                        i += 1
                    _aufgaben_tabelle(doc, [[c.strip() for c in r[2:].split("|")][:3] for r in rows])
                    continue
                if line.startswith("## "):
                    section = line[3:].strip()
                if line.startswith("### "):
                    _heading(doc, line[4:], 4)
                elif line.startswith("## "):
                    _heading(doc, line[3:], 3)
                elif line.startswith("# "):
                    _heading(doc, line[2:], 2)
                elif line.startswith(("- ", "* ")):
                    _bullet(doc, line[2:], briefkopf)
                else:
                    _md_inline(doc.add_paragraph(), line)
            if nur_protokoll:
                buf = io.BytesIO()
                doc.save(buf)
                return buf.getvalue()
            doc.add_page_break()

        _heading(doc, "Transkript", 2)
        for s in t.segments:
            p = doc.add_paragraph()
            r = p.add_run(f"[{fmt_time(s.start)}] ")
            r.font.size = Pt(9)
            if s.speaker:
                r2 = p.add_run(f"{s.speaker}: ")
                r2.bold = True
            p.add_run(s.clean or s.text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
